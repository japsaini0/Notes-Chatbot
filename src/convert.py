from docx import Document
from pptx import Presentation
from pypdf import PdfReader

from src import config


def read_plain(path):
    return path.read_text(encoding="utf-8", errors="ignore")


def read_pdf(path):
    reader = PdfReader(str(path))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def read_docx(path):
    document = Document(str(path))
    return "\n".join(paragraph.text for paragraph in document.paragraphs)


def read_pptx(path):
    presentation = Presentation(str(path))
    lines = []
    for slide in presentation.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                lines.append(shape.text_frame.text)
    return "\n".join(lines)


READERS = {
    ".txt": read_plain,
    ".md": read_plain,
    ".pdf": read_pdf,
    ".docx": read_docx,
    ".pptx": read_pptx,
}


def is_junk(relative_path):
    return any(part.startswith(".") or part == "__MACOSX" for part in relative_path.parts)


def convert_all():
    config.TEXT_DIR.mkdir(parents=True, exist_ok=True)
    for old_file in config.TEXT_DIR.glob("*.txt"):
        old_file.unlink()

    converted = []
    skipped = []

    for path in sorted(config.EXTRACTED_DIR.rglob("*")):
        if not path.is_file():
            continue
        relative = path.relative_to(config.EXTRACTED_DIR)
        if is_junk(relative):
            continue
        reader = READERS.get(path.suffix.lower())
        if reader is None:
            skipped.append(str(relative))
            continue
        try:
            text = reader(path).strip()
        except Exception:
            skipped.append(str(relative))
            continue
        if not text:
            skipped.append(str(relative))
            continue
        output_name = "__".join(relative.parts) + ".txt"
        (config.TEXT_DIR / output_name).write_text(text, encoding="utf-8")
        converted.append(output_name)

    return converted, skipped